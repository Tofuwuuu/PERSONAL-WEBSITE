"""One-page resume tailored for ePLDT — AI Engineer."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Calibri"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

NAME = "Mark Roderick I. Salise"
TITLE = "AI Engineer  |  Full Stack & Pipeline Development"
LINE1 = "Cavite, Philippines  |  0948 253 6598  |  rodericksalise812@gmail.com"
LINE2 = "markroderick.vercel.app  |  github.com/Tofuwuuu"

OUT = r"c:\Users\roder\Downloads\Mark_Salise_Resume_ePLDT_AI_Engineer.docx"
OUT_PROJECT = r"E:\Projects\PERSONAL WEBSITE\Mark_Salise_Resume_ePLDT_AI_Engineer.docx"


def build_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.42)
        s.bottom_margin = Inches(0.42)
        s.left_margin = Inches(0.55)
        s.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = BLACK
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)

    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    return doc


def style_run(run, size, bold=False, italic=False, color=BLACK):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def para(doc, space_after=0.0, space_before=0.0, left=0.0, hanging=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if left:
        p.paragraph_format.left_indent = Inches(left)
    if hanging:
        p.paragraph_format.first_line_indent = Inches(-hanging)
    return p


def header(doc):
    p = para(doc, space_after=0.5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(NAME), 18, bold=True, color=NAVY)

    p = para(doc, space_after=1.5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(TITLE), 9.5, color=NAVY)

    for line in (LINE1, LINE2):
        p = para(doc, space_after=0.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(line), 9, color=BLACK)


def section(doc, title):
    p = para(doc, space_before=5, space_after=2)
    style_run(p.add_run(title.upper()), 10.5, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    bdr.append(bottom)
    pPr.append(bdr)


def entry(doc, left_text, right_text):
    p = para(doc, space_before=3, space_after=0.5)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.35), WD_TAB_ALIGNMENT.RIGHT)
    style_run(p.add_run(left_text), 10, bold=True, color=BLACK)
    if right_text:
        style_run(p.add_run("\t" + right_text), 9, italic=True, color=BLACK)


def bullet(doc, text):
    p = para(doc, space_after=0.5, left=0.16, hanging=0.16)
    style_run(p.add_run("\u2022  " + text), 9.5, color=BLACK)


def meta(doc, text):
    p = para(doc, space_after=0.5, left=0.16)
    style_run(p.add_run(text), 8.5, italic=True, color=BLACK)


def skills(doc, label, items):
    p = para(doc, space_after=1)
    style_run(p.add_run(label + ": "), 9.5, bold=True, color=BLACK)
    style_run(p.add_run(items), 9.5, color=BLACK)


def build():
    doc = build_doc()
    header(doc)

    section(doc, "Professional Summary")
    p = para(doc, space_after=1)
    style_run(
        p.add_run(
            "BS Computer Science graduate with hands-on experience building AI document "
            "processing pipelines, full-stack web applications, and containerized multi-service "
            "systems. Built production-style AI platforms with OCR, NLP entity extraction, "
            "caching, structured API responses, and Docker-based deployment. Experienced in "
            "React, TypeScript, Python, FastAPI, SQL, Redis, and Docker Compose, with strong "
            "foundations in prompt-oriented task design, RAG concepts, and LLM API integration. "
            "Eager to contribute to enterprise AI engineering, pipeline development, and "
            "production system reliability."
        ),
        9.5,
    )

    section(doc, "Technical Skills")
    skills(
        doc,
        "AI & Pipeline",
        "Document AI pipelines, OCR (PyMuPDF/Tesseract), NLP (spaCy NER), prompt engineering, "
        "RAG concepts, LLM API integration, structured inference outputs, caching strategies",
    )
    skills(
        doc,
        "Full Stack",
        "React, TypeScript, JavaScript, Python, FastAPI, Node.js, REST APIs, JSON contracts",
    )
    skills(
        doc,
        "Data & Infrastructure",
        "PostgreSQL, MongoDB, MySQL, Redis, Docker, Docker Compose, Git, Vercel, Railway",
    )
    skills(
        doc,
        "Engineering Practices",
        "API design, technical documentation, environment configuration, testing, debugging, "
        "secure development awareness",
    )
    skills(
        doc,
        "Learning Focus",
        "LangChain, OpenAI/Anthropic APIs, vector databases, RAG pipelines, autoscaling, "
        "cloud-native deployment (Azure/AWS), CI/CD automation",
    )

    section(doc, "Work Experience")

    entry(doc, "AI & Full Stack Engineer  |  Freelance / Self-Employed", "Jan 2024 - Jan 2025")
    bullet(
        doc,
        "Designed and delivered multi-service application architectures connecting React front "
        "ends, Python API services, databases, and containerized environments for paying clients.",
    )
    bullet(
        doc,
        "Built REST API pipelines with structured JSON contracts, audit logging, and role-based "
        "workflow automation across integrated modules.",
    )
    bullet(
        doc,
        "Deployed and maintained Docker-based environments with documented setup for repeatable "
        "delivery and handoff.",
    )
    bullet(
        doc,
        "Developed a Laravel consultation system with automated booking workflows, queued email "
        "delivery, and database notification pipelines.",
    )

    entry(doc, "Web Developer Intern  |  Go Crayons", "Jun 2024 - Aug 2024")
    bullet(
        doc,
        "Resolved client-facing technical issues, coordinated fixes with the team, and validated "
        "completed work before release.",
    )

    section(doc, "Projects")

    entry(doc, "SmartDoc Analyzer — AI Document Processing Pipeline", "2025")
    bullet(
        doc,
        "Built an end-to-end document intelligence pipeline: upload handling, OCR text extraction "
        "(PyMuPDF/Tesseract), spaCy NER, insight generation, SHA-256 deduplication, and Redis "
        "caching with PostgreSQL persistence.",
    )
    bullet(
        doc,
        "Implemented rule-based contract risk analysis (missing clauses, auto-renewal, payment "
        "terms, jurisdiction conflicts) with structured JSON API responses.",
    )
    bullet(
        doc,
        "Containerized all services with Docker Compose for repeatable pipeline execution; "
        "JWT-scoped access for multi-user security.",
    )
    meta(doc, "React  |  TypeScript  |  FastAPI  |  PostgreSQL  |  Redis  |  spaCy  |  Docker  -  Live: smart-doc-analyzer-murex.vercel.app")

    entry(doc, "Hyperledger Alumni Document Verification (Capstone)", "2024 - 2025")
    bullet(
        doc,
        "Architected a distributed verification system with React front end, Node.js Fabric "
        "gateway, FastAPI services, and MongoDB for tamper-evident credential records.",
    )
    meta(doc, "Hyperledger Fabric  |  React  |  FastAPI  |  MongoDB  |  Docker")

    entry(doc, "Government Procurement Workflow Platform", "2024 - 2025")
    bullet(
        doc,
        "Automated multi-stage procurement operations with role-based access, API integrations, "
        "audit logging, and live deployment for client demonstration.",
    )
    meta(doc, "React  |  TypeScript  |  Python  |  MongoDB  |  Docker  -  Live: pams-pied.vercel.app")

    entry(doc, "On-Chain Wallet Risk Analyzer", "2025")
    bullet(
        doc,
        "Built a scoring pipeline with weighted heuristics, provider API clients (Etherscan, "
        "Alchemy), Redis caching, and PostgreSQL report storage with PDF/HTML export.",
    )
    meta(doc, "FastAPI  |  PostgreSQL  |  Redis  |  React  |  TypeScript  -  Live API: api-production-487c.up.railway.app")

    section(doc, "Education")
    entry(doc, "Bachelor of Science in Computer Science", "Graduated September 2025")
    p = para(doc, space_after=0.5)
    style_run(p.add_run("Cavite State University \u2013 Carmona Campus"), 9.5)
    bullet(
        doc,
        "Relevant coursework: Data Structures & Algorithms, Database Systems, Computer Networks, "
        "Operating Systems, Information Assurance & Security, Software Engineering.",
    )
    bullet(
        doc,
        "Treasurer, Young Programmers and Developers' Society (Jun 2022 \u2013 Mar 2024).",
    )

    section(doc, "Additional Information")
    bullet(doc, "Portfolio: markroderick.vercel.app  |  GitHub: github.com/Tofuwuuu")
    bullet(doc, "Available for full-time work; able to start immediately.")
    bullet(doc, "Comfortable with onsite, hybrid, or remote arrangements in Metro Manila / NCR.")

    doc.save(OUT)
    doc.save(OUT_PROJECT)
    print("Saved:", OUT)
    print("Saved:", OUT_PROJECT)


if __name__ == "__main__":
    build()
