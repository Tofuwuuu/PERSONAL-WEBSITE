"""Cover letter (DOCX) for WorkRoute — Customer Service / Customer Experience."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Calibri"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

NAME = "Mark Roderick I. Salise"
TITLE = "Customer Service / Customer Experience Representative  |  Remote"
LINE1 = "Cavite, Philippines  |  0948 253 6598  |  rodericksalise812@gmail.com"
LINE2 = "Languages: English (fluent, spoken & written)  |  Filipino (native)"

OUT = r"c:\Users\roder\Downloads\Mark_Salise_Cover_Letter_Customer_Service.docx"
OUT_PROJECT = r"E:\Projects\PERSONAL WEBSITE\Mark_Salise_Cover_Letter_Customer_Service.docx"


def build_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)

    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    return doc


def style_run(run, size, bold=False, italic=False, color=BLACK):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def para(doc, space_after=0.0, space_before=0.0, left=0.0, hanging=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if left:
        p.paragraph_format.left_indent = Inches(left)
    if hanging:
        p.paragraph_format.first_line_indent = Inches(-hanging)
    return p


def header(doc):
    p = para(doc, space_after=1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(NAME), 18, bold=True, color=NAVY)

    p = para(doc, space_after=1.5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(TITLE), 9.5, color=NAVY)

    for line in (LINE1, LINE2):
        p = para(doc, space_after=0.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(line), 9, color=BLACK)

    p = para(doc, space_before=6, space_after=8)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    bdr.append(bottom)
    pPr.append(bdr)


def body_para(doc, text, space_after=8):
    p = para(doc, space_after=space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style_run(p.add_run(text), 10.5)
    return p


def bullet(doc, text):
    p = para(doc, space_after=2, left=0.22, hanging=0.22)
    style_run(p.add_run("\u2022  " + text), 10.5)


def build():
    doc = build_doc()
    header(doc)

    p = para(doc, space_after=8)
    style_run(p.add_run("Dear WorkRoute Recruitment Team,"), 10.5, bold=True)

    body_para(
        doc,
        "I am writing to express my interest in joining your Customer Experience team. "
        "I am a service-oriented professional with hands-on experience handling client "
        "concerns, resolving issues from first contact through full resolution, and keeping "
        "customers informed with clear, patient communication. I am fluent in English, both "
        "spoken and written, and fully equipped for remote work.",
    )

    body_para(
        doc,
        "During my internship at Go Crayons, I served as a point of contact for client-reported "
        "concerns across several accounts at the same time. My role was to listen carefully, "
        "clarify what the customer actually needed, coordinate the fix with our team, and follow "
        "through until the client confirmed the issue was resolved. Managing multiple open "
        "requests taught me to prioritize under pressure, keep accurate records, and never let a "
        "concern fall through the cracks.",
    )

    body_para(
        doc,
        "In my freelance work, I handled client relationships directly — from the first inquiry "
        "to final delivery and ongoing support. I set realistic expectations, provided regular "
        "updates, and explained technical matters in plain, easy-to-understand language. I also "
        "ran remote walkthrough sessions and wrote simple step-by-step guides so clients could "
        "confidently handle everyday questions on their own. I genuinely enjoy the moment when a "
        "frustrated customer relaxes because someone finally explained things clearly and took "
        "ownership of their problem.",
    )

    body_para(
        doc,
        "What I bring to your team:",
        space_after=4,
    )

    bullet(
        doc,
        "Fluent English communication with a calm, courteous, and professional tone.",
    )
    bullet(
        doc,
        "Proven ability to manage multiple customer concerns at once without losing track of details.",
    )
    bullet(
        doc,
        "Strong technical literacy as a BS Computer Science graduate — an advantage when assisting "
        "customers with accounts, software, or online platform issues.",
    )
    bullet(
        doc,
        "Leadership and accountability from serving as Treasurer of our student developers' "
        "organization, where I managed records and assisted members with their concerns.",
    )
    bullet(
        doc,
        "A complete remote setup: personal Windows computer, headset, stable fiber internet with "
        "mobile backup, and a quiet dedicated workspace.",
    )

    body_para(
        doc,
        "I am open to shifting schedules, including night shift and weekend coverage for "
        "international customers, and I am available to start immediately. I am eager to keep "
        "growing in customer experience and to represent your company well in every interaction "
        "I handle.",
        space_after=6,
    )

    body_para(
        doc,
        "Thank you for taking the time to review my application. I would welcome the opportunity "
        "to discuss how I can contribute to your team, and I am happy to make myself available "
        "for an interview at your convenience.",
    )

    p = para(doc, space_before=6, space_after=1)
    style_run(p.add_run("Sincerely,"), 10.5)

    p = para(doc, space_before=8, space_after=0.5)
    style_run(p.add_run(NAME), 11, bold=True, color=NAVY)

    p = para(doc, space_after=0.5)
    style_run(p.add_run("0948 253 6598  |  rodericksalise812@gmail.com"), 9.5)

    doc.save(OUT)
    doc.save(OUT_PROJECT)
    print("Saved:", OUT)
    print("Saved:", OUT_PROJECT)


if __name__ == "__main__":
    build()
