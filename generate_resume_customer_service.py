"""One-page Customer Service / Customer Experience resume (WorkRoute application)."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Calibri"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

NAME = "Mark Roderick I. Salise"
TITLE = "Customer Service / Customer Experience Representative  |  Remote"
LINE1 = "Cavite, Philippines  |  0948 253 6598  |  rodericksalise812@gmail.com"
LINE2 = "Languages: English (fluent, spoken & written)  |  Filipino (native)"

OUT = r"c:\Users\roder\Downloads\Mark_Salise_CV_Customer_Service.docx"
OUT_PROJECT = r"E:\Projects\PERSONAL WEBSITE\Mark_Salise_CV_Customer_Service.docx"


def build_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.42)
        s.bottom_margin = Inches(0.42)
        s.left_margin = Inches(0.55)
        s.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = BLACK
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)

    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    return doc


def style_run(run, size, bold=False, italic=False, color=BLACK):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def para(doc, space_after=0.0, space_before=0.0, left=0.0, hanging=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if left:
        p.paragraph_format.left_indent = Inches(left)
    if hanging:
        p.paragraph_format.first_line_indent = Inches(-hanging)
    return p


def header(doc):
    p = para(doc, space_after=0.5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(NAME), 18, bold=True, color=NAVY)

    p = para(doc, space_after=1.5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(TITLE), 9.5, color=NAVY)

    for line in (LINE1, LINE2):
        p = para(doc, space_after=0.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(line), 9, color=BLACK)


def section(doc, title):
    p = para(doc, space_before=5, space_after=2)
    style_run(p.add_run(title.upper()), 10.5, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    bdr.append(bottom)
    pPr.append(bdr)


def entry(doc, left_text, right_text):
    p = para(doc, space_before=3, space_after=0.5)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.35), WD_TAB_ALIGNMENT.RIGHT)
    style_run(p.add_run(left_text), 10, bold=True, color=BLACK)
    if right_text:
        style_run(p.add_run("\t" + right_text), 9, italic=True, color=BLACK)


def bullet(doc, text):
    p = para(doc, space_after=0.5, left=0.16, hanging=0.16)
    style_run(p.add_run("\u2022  " + text), 9.5, color=BLACK)


def skills(doc, label, items):
    p = para(doc, space_after=1)
    style_run(p.add_run(label + ": "), 9.5, bold=True, color=BLACK)
    style_run(p.add_run(items), 9.5, color=BLACK)


def build():
    doc = build_doc()
    header(doc)

    section(doc, "Professional Summary")
    p = para(doc, space_after=1)
    style_run(
        p.add_run(
            "Customer-focused support professional with hands-on experience handling "
            "client concerns, resolving issues, and maintaining clear communication across "
            "multiple accounts. Skilled at listening to the customer's problem, explaining "
            "solutions in simple language, following through until the concern is fully "
            "resolved, and keeping organized records of every interaction. Fluent in English "
            "(spoken and written), patient under pressure, and comfortable in fast-paced "
            "remote environments. BS Computer Science graduate with strong technical literacy, "
            "making it easy to support customers on software, accounts, and online platforms."
        ),
        9.5,
    )

    section(doc, "Core Strengths")
    skills(
        doc,
        "Customer Service",
        "Customer inquiry handling, complaint resolution, de-escalation, follow-through to "
        "closure, expectation setting, empathetic and professional tone",
    )
    skills(
        doc,
        "Communication",
        "Fluent English (spoken & written), clear email and chat etiquette, simplifying "
        "technical topics for non-technical customers, active listening, status updates",
    )
    skills(
        doc,
        "Organization",
        "Managing multiple concurrent requests, prioritization, accurate documentation, "
        "record keeping, attention to detail, meeting turnaround commitments",
    )
    skills(
        doc,
        "Technical Support",
        "Account and access issues, software troubleshooting, browser and device basics, "
        "escalation with clear reproduction steps, ticket tracking and organization",
    )
    skills(
        doc,
        "Tools",
        "Microsoft 365 (Outlook, Word, Excel), Google Workspace (Gmail, Docs, Sheets), "
        "Zoom / Google Meet, Facebook & Messenger business tools, CRM and ticketing workflows",
    )
    skills(
        doc,
        "Personal Qualities",
        "Reliable, patient, self-driven, coachable, calm with upset customers, team player, "
        "leadership potential",
    )

    section(doc, "Work Experience")

    entry(doc, "Client Support & Web Developer Intern  |  Go Crayons", "Jun 2024 - Aug 2024")
    bullet(
        doc,
        "Served as a point of contact for client-reported concerns — listened to the issue, "
        "clarified details, and provided clear updates until each concern was resolved.",
    )
    bullet(
        doc,
        "Handled multiple client accounts at the same time, prioritizing urgent requests "
        "while keeping every open item tracked and accounted for.",
    )
    bullet(
        doc,
        "Coordinated with internal team members to resolve issues quickly, then verified the "
        "fix before confirming completion back to the client.",
    )
    bullet(
        doc,
        "Maintained professional, courteous communication in a service-oriented environment "
        "with concurrent requests and shifting priorities.",
    )

    entry(doc, "Client Relations & Technical Support  |  Freelance / Self-Employed", "Jan 2024 - Jan 2025")
    bullet(
        doc,
        "Managed client relationships directly from first inquiry through delivery — "
        "gathering requirements, setting realistic expectations, and giving regular updates.",
    )
    bullet(
        doc,
        "Responded to client questions and reported problems, troubleshooting issues and "
        "explaining the cause and solution in plain, non-technical language.",
    )
    bullet(
        doc,
        "Conducted remote walkthrough and training sessions so clients and their staff could "
        "confidently use the systems delivered to them.",
    )
    bullet(
        doc,
        "Created step-by-step guides and instructions that reduced repeat questions and "
        "helped clients solve simple concerns on their own.",
    )
    bullet(
        doc,
        "Kept organized notes of every request, change, and resolution for accurate follow-up.",
    )

    section(doc, "Leadership & Additional Experience")

    entry(doc, "Treasurer  |  Young Programmers and Developers' Society", "Jun 2022 - Mar 2024")
    bullet(
        doc,
        "Handled organization funds and financial records with complete accuracy and "
        "accountability, reporting regularly to officers and members.",
    )
    bullet(
        doc,
        "Assisted members with questions and concerns, coordinated events, and communicated "
        "clearly across a large group — building patience and people-handling skills.",
    )

    entry(doc, "Capstone Project — Stakeholder Coordination", "2024 - 2025")
    bullet(
        doc,
        "Worked with university staff and stakeholders to understand their needs, presented "
        "progress in plain language, and adjusted based on their feedback.",
    )

    section(doc, "Education")
    entry(doc, "Bachelor of Science in Computer Science", "Graduated September 2025")
    p = para(doc, space_after=0.5)
    style_run(p.add_run("Cavite State University - Carmona Campus, Philippines"), 9.5)
    bullet(
        doc,
        "Coursework included communication, technical writing, database systems, computer "
        "networks, and information security — supporting strong technical support ability.",
    )

    section(doc, "Remote Work Setup & Availability")
    bullet(doc, "Own Windows 10/11 computer with headset and webcam; quiet dedicated workspace.")
    bullet(doc, "Stable home fiber internet with mobile data as backup connection.")
    bullet(
        doc,
        "Open to shifting schedules, including night shift and weekend coverage for "
        "international customers.",
    )
    bullet(doc, "Available to start immediately; comfortable with 100% remote work.")

    doc.save(OUT)
    doc.save(OUT_PROJECT)
    print("Saved:", OUT)
    print("Saved:", OUT_PROJECT)


if __name__ == "__main__":
    build()
