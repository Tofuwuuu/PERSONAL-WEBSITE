"""One-page Software Engineer resume with school + work experience."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Calibri"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

NAME = "Mark Roderick I. Salise"
TITLE = "Software Engineer"
LINE1 = "Cavite, Philippines  |  0948 253 6598  |  rodericksalise812@gmail.com"
LINE2 = "markroderick.vercel.app  |  github.com/Tofuwuuu"

OUT = r"c:\Users\roder\Downloads\Mark_Salise_Resume_Software_Engineer.docx"
OUT_PROJECT = r"E:\Projects\PERSONAL WEBSITE\Mark_Salise_Resume_Software_Engineer.docx"


def build_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.42)
        s.bottom_margin = Inches(0.42)
        s.left_margin = Inches(0.55)
        s.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = BLACK
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)

    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    return doc


def style_run(run, size, bold=False, italic=False, color=BLACK):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def para(doc, space_after=0.0, space_before=0.0, left=0.0, hanging=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if left:
        p.paragraph_format.left_indent = Inches(left)
    if hanging:
        p.paragraph_format.first_line_indent = Inches(-hanging)
    return p


def header(doc):
    p = para(doc, space_after=0.5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(NAME), 18, bold=True, color=NAVY)

    p = para(doc, space_after=1.5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(TITLE), 9.5, color=NAVY)

    for line in (LINE1, LINE2):
        p = para(doc, space_after=0.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(line), 9, color=BLACK)


def section(doc, title):
    p = para(doc, space_before=5, space_after=2)
    style_run(p.add_run(title.upper()), 10.5, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    bdr.append(bottom)
    pPr.append(bdr)


def entry(doc, left_text, right_text):
    p = para(doc, space_before=3, space_after=0.5)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.35), WD_TAB_ALIGNMENT.RIGHT)
    style_run(p.add_run(left_text), 10, bold=True, color=BLACK)
    if right_text:
        style_run(p.add_run("\t" + right_text), 9, italic=True, color=BLACK)


def bullet(doc, text):
    p = para(doc, space_after=0.5, left=0.16, hanging=0.16)
    style_run(p.add_run("\u2022  " + text), 9.5, color=BLACK)


def meta(doc, text):
    p = para(doc, space_after=0.5, left=0.16)
    style_run(p.add_run(text), 8.5, italic=True, color=BLACK)


def skills(doc, label, items):
    p = para(doc, space_after=1)
    style_run(p.add_run(label + ": "), 9.5, bold=True, color=BLACK)
    style_run(p.add_run(items), 9.5, color=BLACK)


def build():
    doc = build_doc()
    header(doc)

    section(doc, "Professional Summary")
    p = para(doc, space_after=1)
    style_run(
        p.add_run(
            "Software Engineer and BS Computer Science graduate with hands-on experience "
            "shipping full-stack web applications, REST APIs, databases, and containerized "
            "deployments. Built production-style systems for freelance clients and academic "
            "stakeholders using React, TypeScript, Python, FastAPI, Laravel, SQL, and Docker. "
            "Strong foundation from Computer Science coursework in networks, operating systems, "
            "databases, and information security, with leadership experience as organization "
            "treasurer. Eager to contribute as a delivery-focused software engineer."
        ),
        9.5,
    )

    section(doc, "Technical Skills")
    skills(
        doc,
        "Languages",
        "TypeScript, JavaScript, Python, PHP, SQL, HTML, CSS, Go (chaincode basics)",
    )
    skills(
        doc,
        "Frontend",
        "React, Vite, responsive UI, REST API integration, form/validation flows",
    )
    skills(
        doc,
        "Backend & Data",
        "FastAPI, Laravel, Node.js gateway patterns, PostgreSQL, MySQL, MongoDB, Redis",
    )
    skills(
        doc,
        "Tools & Practices",
        "Git, GitHub, Docker, Docker Compose, Vercel, Railway, technical documentation",
    )
    skills(
        doc,
        "Domains",
        "Full-stack delivery, workflow automation, AI document pipelines, blockchain verification",
    )

    section(doc, "Work Experience")

    entry(doc, "Software Engineer  |  Freelance / Self-Employed", "Jan 2024 - Jan 2025")
    bullet(
        doc,
        "Delivered a multi-stage procurement platform for a paying client with role-based "
        "workflows, REST APIs, audit logging, and Docker-based deployment.",
    )
    bullet(
        doc,
        "Built React/TypeScript front ends connected to Python FastAPI services and MongoDB "
        "for operational modules (purchase requests through inventory).",
    )
    bullet(
        doc,
        "Developed a Laravel 11 consultation booking system with validation rules, queued "
        "email delivery, MySQL persistence, and Docker setup.",
    )
    bullet(
        doc,
        "Documented APIs, environment setup, and handoff notes so stakeholders could run "
        "and maintain delivered systems.",
    )

    entry(doc, "Web Developer Intern  |  Go Crayons", "Jun 2024 - Aug 2024")
    bullet(
        doc,
        "Supported client websites with updates, issue resolution, and quality checks across "
        "multiple accounts.",
    )
    bullet(
        doc,
        "Communicated clearly with teammates and clients while coordinating fixes and "
        "validating completed work before release.",
    )

    section(doc, "Projects")

    entry(doc, "Hyperledger Alumni Document Verification (Capstone Thesis)", "2024 - 2025")
    bullet(
        doc,
        "Capstone for Cavite State University – Carmona: permissioned Hyperledger Fabric "
        "network for tamper-evident alumni credential verification.",
    )
    bullet(
        doc,
        "Built React front end, FastAPI services, MongoDB storage, and Node.js Fabric gateway; "
        "delivered demo and documentation for stakeholders.",
    )
    meta(doc, "Hyperledger Fabric  |  React  |  FastAPI  |  MongoDB  |  Docker  -  Live: hyperledger-document-verification.vercel.app")

    entry(doc, "SmartDoc Analyzer — AI Document Platform", "2025")
    bullet(
        doc,
        "Built an AI document pipeline with OCR, spaCy NER, Redis caching, PostgreSQL storage, "
        "and structured API responses for a React/TypeScript UI.",
    )
    meta(doc, "React  |  TypeScript  |  FastAPI  |  PostgreSQL  |  Redis  |  Docker  -  Live: smart-doc-analyzer-murex.vercel.app")

    entry(doc, "Government Procurement Workflow Platform", "2024 - 2025")
    bullet(
        doc,
        "Automated procurement stages with role-based access, API integrations, and live "
        "deployment for client demonstration and handoff.",
    )
    meta(doc, "React  |  TypeScript  |  Python  |  MongoDB  |  Docker  -  Live: pams-pied.vercel.app")

    section(doc, "Education")
    entry(doc, "Bachelor of Science in Computer Science", "Graduated September 2025")
    p = para(doc, space_after=0.5)
    style_run(p.add_run("Cavite State University – Carmona Campus"), 9.5)
    bullet(
        doc,
        "Relevant coursework: Computer Programming, Data Structures & Algorithms, Database "
        "Systems, Computer Networks, Operating Systems, Information Assurance & Security, "
        "Software Engineering, Web Development.",
    )
    bullet(
        doc,
        "Capstone Thesis: Hyperledger Fabric alumni document verification system — designed, "
        "implemented, and defended end-to-end for university stakeholders.",
    )
    bullet(
        doc,
        "Academic projects: configured local networks and development environments; built "
        "full-stack applications integrating front end, APIs, and databases for course "
        "requirements and laboratory work.",
    )
    bullet(
        doc,
        "Treasurer, Young Programmers and Developers' Society (Jun 2022 – Mar 2024) — "
        "organization finances, record keeping, event coordination, and team communication.",
    )

    section(doc, "Additional Information")
    bullet(doc, "Available for full-time software engineering roles; able to start immediately.")
    bullet(doc, "Portfolio: markroderick.vercel.app  |  GitHub: github.com/Tofuwuuu")
    bullet(doc, "Comfortable with onsite, hybrid, or remote arrangements in Metro Manila / NCR.")

    doc.save(OUT)
    doc.save(OUT_PROJECT)
    print("Saved:", OUT)
    print("Saved:", OUT_PROJECT)


if __name__ == "__main__":
    build()
